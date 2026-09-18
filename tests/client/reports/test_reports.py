"""PDF and Excel generation: render real documents to temp files and read them back."""

import os
import subprocess
from datetime import datetime

import pytest
from pypdf import PdfReader
import openpyxl

from GlobalData import globalData
from models.PTW import PTW
from network.clientRequests import ClientRequests
from reports import ReportGenerator as rg_module
from reports.ReportGenerator import ReportGenerator
from ptw_factory import ptw_payload, full_chain, running_cycle


@pytest.fixture
def opened(monkeypatch):
    """Capture every file the generator would open in a viewer, instead of opening it."""
    files = []
    monkeypatch.setattr(ReportGenerator, 'openPDF', staticmethod(lambda path: files.append(path)))
    monkeypatch.setattr(rg_module.subprocess, 'call', lambda args, *a, **k: files.append(args[-1]))
    monkeypatch.setattr(rg_module.os, 'startfile', lambda path: files.append(path), raising=False)
    yield files
    for f in files:
        try:
            os.remove(f)
        except OSError:
            pass


@pytest.fixture
def offline_requests(monkeypatch):
    monkeypatch.setattr(ClientRequests, 'getPTWSpecificRiskAssessment', staticmethod(lambda user, pid: (None, None)))
    monkeypatch.setattr(ClientRequests, 'getPtwAttachmentNames', staticmethod(lambda user, pid: (None, [])))
    monkeypatch.setattr(ClientRequests, 'getMIWI', staticmethod(lambda *a, **k: ('offline', None)))


def text_of(path):
    return '\n'.join((page.extract_text() or '') for page in PdfReader(path).pages)


class TestPtwReport:
    def test_running_permit_report(self, qapp, known_users, offline_requests, opened):
        ptw = PTW(ptw_payload(id=42, type='Hot', equipment='K-201', description='Weld the bracket', hazards=['Noise'],
                              controls=['Hearing Protection', 'Initial Gas Test'], tools=['Hand Tools'],
                              approvals=full_chain('Hot'), run_cycles=[running_cycle(datetime(2026, 9, 1, 8, 0), ia='issuing')]))
        err = ReportGenerator.ptwReport(known_users['coord'], ptw)
        assert err is None
        assert len(opened) >= 2, opened                  # the permit itself + the merged required docs
        permit = text_of(opened[0])
        for needle in ('42', 'Hot', 'K-201', 'Weld the bracket', 'Turbo', 'Running 08:00 - 19:00', 'User Turbo', 'Noise', 'Hearing Protection'):
            assert needle in permit, needle
        # signature blocks name the whole chain incl. PGM/DFGM for hot work
        for role in ('Coordinator', 'Issuing', 'HSE Engineer', 'PGM', 'DFGM'):
            assert role in permit, role
        assert 'Page 1 of' in permit

    def test_required_docs_are_merged_and_normalised_to_portrait_a4(self, qapp, known_users, offline_requests, opened):
        ptw = PTW(ptw_payload(id=7, type='Hot', controls=['Initial Gas Test'], hazards=['Working at Height'], linked_ics=['3']))
        ReportGenerator.ptwReport(known_users['user_turbo'], ptw)
        merged = [p for p in opened if 'required-docs' in os.path.basename(p)]
        assert len(merged) == 1
        reader = PdfReader(merged[0])
        expected_keys = ptw.requiredDocsToPrint()
        assert {'toolbox', 'audit', 'gas-test', 'swc-hot-work', 'swc-working-at-height', 'swc-energy-isolation', 'swc-de-isolation'} <= set(expected_keys)
        assert len(reader.pages) >= len(expected_keys)
        for i, page in enumerate(reader.pages):
            w, h = float(page.mediabox.width), float(page.mediabox.height)
            assert abs(w - 595.28) < 1 and abs(h - 841.89) < 1, f"page {i} is {w}x{h}, not portrait A4"
            assert page.rotation == 0

    def test_arabic_description_renders(self, qapp, known_users, offline_requests, opened):
        ptw = PTW(ptw_payload(id=9, description='استبدال حشية المضخة الرئيسية', equipment='مضخة-١٠١'))
        assert ReportGenerator.ptwReport(known_users['user_turbo'], ptw) is None
        reader = PdfReader(opened[0])
        assert len(reader.pages) >= 2
        fonts = set()
        for page in reader.pages:
            res = page.get('/Resources') or {}
            for f in (res.get('/Font') or {}).values():
                fonts.add(str(f.get_object().get('/BaseFont')))
        assert any('Naskh' in f or 'Arabic' in f for f in fonts), fonts    # the bundled Arabic font was actually used

    def test_unknown_users_do_not_crash_the_report(self, qapp, offline_requests, opened):
        # nobody in globalData.allUsers: every name falls back gracefully
        from conftest import make_user
        ptw = PTW(ptw_payload(id=3, approvals=full_chain('Cold')))
        assert ReportGenerator.ptwReport(make_user(), ptw) is None
        assert 'None' in text_of(opened[0])            # the documented fallback for an unknown requestor


class TestExcelExport:
    def test_export_ptws(self, qapp, known_users, opened):
        ptws = [
            PTW(ptw_payload(id=1, description='first')),
            PTW(ptw_payload(id=2, type='Hot', approvals=full_chain('Hot'), run_cycles=[running_cycle(datetime(2026, 9, 1, 8, 0))],
                            description='وصف بالعربية')),
        ]
        ReportGenerator.exportPTWs(ptws)
        assert len(opened) == 1 and opened[0].endswith('.xlsx')
        ws = openpyxl.load_workbook(opened[0]).active
        header = [c.value for c in ws[1]]
        assert header == ['PTW#', 'Type', 'Status', 'Date', 'Department', 'Requestor', 'PA', 'Location', 'Area Class', 'Equipment', 'Description']
        rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
        assert len(rows) == 2
        assert rows[0][0] == 1 and rows[0][2] == 'Under Review'
        assert rows[1][1] == 'Hot' and rows[1][2] == 'Running 08:00 - 19:00'
        assert rows[1][6] == 'User Turbo'                          # PA resolved to a display name
        assert ws.freeze_panes == 'A2'
        # type colouring is carried into the cells: Cold and Hot rows get different solid fills
        cold, hot = ws.cell(row=2, column=1).fill, ws.cell(row=3, column=1).fill
        assert cold.fill_type == 'solid' and hot.fill_type == 'solid'
        assert cold.fgColor.rgb != hot.fgColor.rgb
