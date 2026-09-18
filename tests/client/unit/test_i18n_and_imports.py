"""Translations, RTL detection, tabular import parsing, and the bulk user import."""

import csv
import os

import pytest
import openpyxl

import helper.i18n as i18n
from helper.i18n import t
from helper.utils import parseTabularFile
from reports.ImportUsersExcel import ImportUsersExcel, HEADERS

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
TEST_DATA = os.path.join(ROOT, 'test_data')


@pytest.fixture(autouse=True)
def restore_english():
    yield
    i18n.init('en')


class TestI18n:
    def test_english_is_identity(self):
        i18n.init('en')
        assert t('Approved PTWs') == 'Approved PTWs'
        assert i18n.is_rtl() is False and i18n.current_lang() == 'en'

    def test_arabic_file_loads_and_translates(self):
        i18n.init('ar')
        assert i18n.is_rtl() is True
        assert i18n._translations, "client/translations/ar.json did not load (see the path comment in helper/i18n.py)"
        translated = [k for k in ('View', 'Edit', 'Delete', 'Print', 'Approved PTWs') if t(k) != k]
        assert translated, "none of the common UI strings have an Arabic translation"
        assert t('some string nobody translated') == 'some string nobody translated'

    def test_arabic_file_is_well_formed(self):
        import json
        with open(os.path.join(ROOT, 'client', 'translations', 'ar.json'), encoding='utf-8') as f:
            data = json.load(f)
        assert all(isinstance(k, str) and isinstance(v, str) for k, v in data.items())
        empty = [k for k, v in data.items() if not v.strip()]
        assert not empty, f"empty translations: {empty[:5]}"

    def test_unknown_language_falls_back_silently(self):
        i18n.init('xx')
        assert t('Edit') == 'Edit' and i18n.is_rtl() is False


class TestParseTabularFile:
    def test_csv_columns_are_reordered_and_matched_loosely(self, tmp_path):
        p = tmp_path / 'u.csv'
        with open(p, 'w', newline='', encoding='utf-8-sig') as f:
            w = csv.writer(f)
            w.writerow(['  NAME ', 'username', 'Role\n', 'Department', 'email', 'ext', 'ignored'])
            w.writerow(['Bob', 'bob', 'User', 'Turbo', 'b@x', '', 'zzz'])
            w.writerow(['', '', '', '', '', '', ''])
        rows = parseTabularFile(str(p), HEADERS)
        assert rows[0] == ['bob', 'Bob', 'User', 'Turbo', 'b@x', '']
        assert rows[1] == [''] * 6            # blank rows are kept; callers filter

    def test_xlsx(self, tmp_path):
        p = tmp_path / 'u.xlsx'
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(HEADERS)
        ws.append(['ann', 'Ann', 'Isolator', 'Prod', None, 123])
        wb.save(p)
        assert parseTabularFile(str(p), HEADERS) == [['ann', 'Ann', 'Isolator', 'Prod', '', '123']]

    def test_missing_column_is_reported(self, tmp_path):
        p = tmp_path / 'u.csv'
        p.write_text('Username,Name\nx,y\n')
        with pytest.raises(ValueError, match='Missing required column'):
            parseTabularFile(str(p), HEADERS)

    def test_empty_file(self, tmp_path):
        p = tmp_path / 'e.csv'
        p.write_text('')
        assert parseTabularFile(str(p), HEADERS) == []


class TestImportUsers:
    @pytest.mark.parametrize("filename", ['test_users_import.csv', 'test_users_import.xlsx'])
    def test_sample_files_parse_identically(self, filename):
        rows = ImportUsersExcel.parseFile(os.path.join(TEST_DATA, filename), existingUsernames={'admin'})
        by_name = {r.username: r for r in rows}
        assert by_name['u'].status == 'Ready to import'
        assert by_name['u'].user.getRole() == 'User' and by_name['u'].user.getDepartment() == 'Turbo'
        assert by_name['u'].user.getPassword()                      # generated
        assert by_name['admin'].status == "Skipped: Username 'admin' already exists"
        assert by_name['g'].user.getRole() == 'HSE Engineer'
        assert by_name['t'].user.getRole() == 'Gas Tester'
        assert rows[0].rowNum == 2                                    # spreadsheet row numbers, header is row 1
        assert len({r.user.getPassword() for r in rows if r.user}) == len([r for r in rows if r.user])   # unique passwords

    def test_validation_messages(self, tmp_path):
        p = tmp_path / 'u.csv'
        p.write_text('Username,Name,Role,Department,Email,EXT\n'
                     ',No Name,User,Turbo,,\n'
                     'dup,A,user,turbo,,\n'
                     'dup,B,User,Turbo,,\n'
                     'nn,,User,Turbo,,\n'
                     'br,Bad Role,Wizard,Turbo,,\n'
                     'bd,Bad Dept,User,Narnia,,\n'
                     ',,,,,\n')
        rows = ImportUsersExcel.parseFile(str(p), existingUsernames=set())
        assert [r.status for r in rows] == [
            "Skipped: Username can't be empty",
            'Ready to import',
            "Skipped: Username 'dup' already exists",
            "Skipped: Name can't be empty",
            "Skipped: Invalid role 'Wizard'",
            "Skipped: Invalid department 'Narnia'",
        ]
        assert rows[1].user.getRole() == 'User' and rows[1].user.getDepartment() == 'Turbo'   # case-normalised

    def test_export_result_round_trip(self, tmp_path):
        rows = ImportUsersExcel.parseFile(os.path.join(TEST_DATA, 'test_users_import.csv'), existingUsernames=set())
        out = tmp_path / 'result.xlsx'
        ImportUsersExcel.exportResult(str(out), rows)
        ws = openpyxl.load_workbook(out).active
        header = [c.value for c in ws[1]]
        assert header == HEADERS + ['Password', 'Status']
        assert ws.max_row == len(rows) + 1
